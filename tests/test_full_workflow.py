"""Behavioral checks of the integrated pipeline, using explicitly synthetic inputs."""
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import corpus
import parse_problem
import pdf_audit
import prose_lint
import render_ai_usage
import render_paper
import workflow
import pymupdf


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        workflow.init(self.root, "cumcm", 2026)
        (self.root / "proof.txt").write_text("actual fixture evidence", encoding="utf-8")
        self.receipt = {"status": "passed", "checks": ["fixture checked"], "artifacts": ["proof.txt"], "issues": []}

    def test_resume_preserves_state_and_rejects_different_competition(self):
        workflow.complete(self.root, 0, self.receipt)
        self.assertEqual(workflow.init(self.root, "cumcm", 2026)["current_stage"], 1)
        with self.assertRaises(ValueError):
            workflow.init(self.root, "mcm", 2027)

    def test_out_of_order_completion_is_rejected(self):
        with self.assertRaises(ValueError):
            workflow.complete(self.root, 2, self.receipt)

    def test_missing_artifact_and_high_issue_block(self):
        bad = {**self.receipt, "artifacts": ["missing"]}
        with self.assertRaises(ValueError):
            workflow.complete(self.root, 0, bad)
        with self.assertRaises(ValueError):
            workflow.complete(self.root, 0, {**self.receipt, "issues": [{"severity": "high"}]})

    def test_changed_upstream_cannot_silently_advance(self):
        workflow.complete(self.root, 0, self.receipt)
        (self.root / "proof.txt").write_text("changed", encoding="utf-8")
        self.assertEqual(workflow.status(self.root)["changed_artifacts"], ["proof.txt"])
        with self.assertRaisesRegex(ValueError, "Upstream"):
            workflow.complete(self.root, 1, self.receipt)
        workflow.rollback(self.root, 0, "recheck source")
        self.assertEqual(workflow.status(self.root)["stale_stages"], [0])
        workflow.complete(self.root, 0, self.receipt)
        self.assertEqual(workflow.status(self.root)["changed_artifacts"], [])

    def test_formal_core_stage_requires_actual_human_review(self):
        state = workflow.load(self.root)
        state["workflow"]["formal_contest"] = True
        workflow.save(self.root / "state/decision_log.json", state)
        workflow.complete(self.root, 0, self.receipt)
        with self.assertRaisesRegex(ValueError, "human review"):
            workflow.complete(self.root, 1, self.receipt)

    def test_ten_stages_finish_only_with_compliance_evidence(self):
        for stage in range(9):
            workflow.complete(self.root, stage, self.receipt)
        with self.assertRaisesRegex(ValueError, "compliance"):
            workflow.complete(self.root, 9, self.receipt)
        state = workflow.load(self.root)
        state["stages"]["9"]["compliance_checks"] = {k: True for k in state["stages"]["9"]["compliance_checks"]}
        state["compliance"]["ai_usage"] = []  # synthetic fixture, not a real no-use declaration
        workflow.save(self.root / "state/decision_log.json", state)
        (self.root / "state/final-gate.json").write_text(
            json.dumps({"status": "READY", "inputs": {}}), encoding="utf-8"
        )  # synthetic workflow-only fixture; final_gate.py has separate integration tests
        workflow.complete(self.root, 9, self.receipt)
        self.assertTrue(workflow.status(self.root)["finished"])


class ParsingTests(unittest.TestCase):
    def test_questions_source_positions_and_csv(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            p = root / "problem.md"
            p.write_text("背景资料\n问题一：预测需求。\n保留时间顺序。\n问题二：最小化成本。\n给出约束。", encoding="utf-8")
            a = root / "data.csv"
            a.write_text("time,value\n1,2\n2,3\n", encoding="utf-8-sig")
            result = parse_problem.parse(p, [a])
            self.assertEqual(result["candidate_count"], 2)
            self.assertIsNone(result["question_count"])
            self.assertEqual(result["questions"][0]["source"]["line"], 2)
            self.assertEqual(result["questions"][1]["type_candidates"][0]["type"], "optimization")
            self.assertEqual(result["attachments"][0]["rows"], 2)

    def test_docx_and_excel_are_read(self):
        from docx import Document
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            d = Document(); d.add_paragraph("问题1：拟合参数并解释。")
            d.save(root / "q.docx")
            wb = openpyxl.Workbook(); ws = wb.active; ws.append(["x", "y"]); ws.append([1, 2]); wb.save(root / "a.xlsx")
            result = parse_problem.parse(root / "q.docx", [root / "a.xlsx"])
            self.assertEqual(result["candidate_count"], 1)
            self.assertEqual(result["attachments"][0]["rows"], 1)

    def test_image_only_pdf_is_not_successfully_understood(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); pdf = root / "scan.pdf"
            doc = pymupdf.open(); doc.new_page(); doc.save(pdf); doc.close()
            result = parse_problem.parse(pdf, [])
            self.assertIn("image_only", result["source_warnings"])
            self.assertEqual(result["candidate_count"], 0)


class PdfCorpusTests(unittest.TestCase):
    @staticmethod
    def make_pdf(path, page_texts):
        doc = pymupdf.open()
        for text in page_texts:
            p = doc.new_page(); p.insert_text((72, 72), text)
        doc.save(path); doc.close()

    def test_cumcm_counts_body_not_appendix(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); pdf = root / "paper.pdf"
            self.make_pdf(pdf, ["abstract", "body 1", "body 2", "Appendix", "code"])
            r = pdf_audit.audit(pdf, "cumcm", body_start=2)
            self.assertEqual((r["counted_start"], r["counted_end"], r["counted_pages"]), (2, 3, 2))

    def test_shared_appendix_page_does_not_undercount_body(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); pdf = root / "paper.pdf"
            self.make_pdf(pdf, ["abstract", "body 1", "final paragraph\nAppendix", "code"])
            r = pdf_audit.audit(pdf, "cumcm", body_start=2)
            self.assertEqual(r["counted_end"], 3)
            self.assertTrue(any(i["code"] == "shared-boundary" for i in r["issues"]))

    def test_mcm_counts_appendix_but_excludes_ai_report(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); pdf = root / "paper.pdf"
            self.make_pdf(pdf, ["summary", "body", "Appendix", "code", "Report on Use of AI", "AI detail"])
            r = pdf_audit.audit(pdf, "mcm")
            self.assertEqual(r["counted_pages"], 4)

    def test_render_page_pngs(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); pdf = root / "paper.pdf"
            self.make_pdf(pdf, ["abstract", "body", "Appendix"])
            out = root / "pages"
            pdf_audit.audit(pdf, "cumcm", body_start=2, render_dir=out)
            self.assertEqual(len(list(out.glob("*.png"))), 3)

    def test_corpus_deduplicates_and_excludes_scans(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); pdir = root / "papers"; pdir.mkdir()
            self.make_pdf(pdir / "a.pdf", ["2024 problem A " + "model result " * 30])
            (pdir / "b.pdf").write_bytes((pdir / "a.pdf").read_bytes())
            self.make_pdf(pdir / "scan.pdf", [""])
            idx = corpus.ingest(pdir)
            self.assertEqual(len(idx["papers"]), 1)
            self.assertGreaterEqual(len(idx["excluded"]), 2)


class DisclosureTests(unittest.TestCase):
    def test_used_ai_declaration_precedes_references(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); state = ROOT / "tests/fixtures/ai_usage_valid.json"
            paper = root / "paper"; support = root / "support"
            render_ai_usage.render_reports(state, paper, support, "cumcm")
            decl = (paper / "AI工具使用声明.md").read_text(encoding="utf-8")
            self.assertIn("参考文献之前", decl)
            self.assertTrue((support / "AI工具使用详情.pdf").is_file())


class StyleTests(unittest.TestCase):
    def test_style_only_edits_preserve_protected_tokens(self):
        before = "MAPE=8.2%，见式(3)，Smith et al. [1]，且不超过20。"
        after = "MAPE=8.2%。根据式(3)与Smith et al. [1]，该值不超过20。"
        self.assertEqual(prose_lint.protected(before), prose_lint.protected(after))

    def test_protected_values_cannot_be_changed_by_polishing(self):
        before = "RMSE=1.25，p<0.05，不超过30。"
        after = "RMSE=1.20，p<0.05，不超过30。"
        self.assertNotEqual(prose_lint.protected(before), prose_lint.protected(after))

    def test_code_and_equations_are_not_style_linted(self):
        text = "这是一个需要进行分析的句子。\n```python\nx = 1  # 需要进行保留\n```\n$$y = x + 1$$"
        issues = prose_lint.lint(text, "zh")
        self.assertTrue(any(i["line"] == 1 for i in issues))
        self.assertFalse(any(i["line"] in {3, 5} for i in issues))
