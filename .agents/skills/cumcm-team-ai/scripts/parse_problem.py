#!/usr/bin/env python3
"""Extract source-located problem candidates; semantic review belongs to the agent."""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

TYPES = {
    "optimization": ["最优", "最大", "最小", "规划", "调度", "optimal", "maximize", "minimize"],
    "prediction": ["预测", "时间序列", "forecast", "predict"],
    "classification": ["分类", "识别", "判别", "classify", "classification"],
    "evaluation": ["评价", "排名", "权重", "rank", "evaluate"],
    "simulation": ["模拟", "仿真", "微分", "动力学", "simulate", "differential"],
    "statistics": ["回归", "相关", "检验", "regression", "correlation"],
}
QUESTION = re.compile(r"^\s*(?:#{1,6}\s*)?(?:问题\s*([一二三四五六七八九十\d]+)|(?:Question|Problem|Task)\s+(\d+)|[（(](\d+)[）)]|([1-9]\d*)[.、](?=\s|[^\d]))", re.I)


def extract_pages(path: Path) -> tuple[list[str], list[str]]:
    suffix = path.suffix.lower()
    warnings = []
    if suffix == ".pdf":
        import pymupdf
        with pymupdf.open(path) as doc:
            pages = [page.get_text(sort=True) for page in doc]
        for n, text in enumerate(pages, 1):
            if len(text.strip()) < 30:
                warnings.append(f"page {n}: image-only/sparse; OCR or visual inspection required")
    elif suffix == ".docx":
        from docx import Document
        doc = Document(path)
        pages = ["\n".join([p.text for p in doc.paragraphs] +
                           [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows])]
        warnings.append("DOCX: logical text blocks; page numbers and equation objects require visual inspection")
    elif suffix in {".md", ".txt"}:
        pages = [path.read_text(encoding="utf-8-sig")]
    else:
        raise ValueError(f"Unsupported problem format: {suffix}")
    return pages, warnings


def classify(text: str) -> list[dict]:
    low = text.lower()
    found = [{"type": kind, "evidence_keywords": [k for k in keys if k in low]}
             for kind, keys in TYPES.items()]
    return sorted([x for x in found if x["evidence_keywords"]],
                  key=lambda x: -len(x["evidence_keywords"]))


def attachment_info(path: Path) -> dict:
    result = {"path": str(path.resolve()), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
              "bytes": path.stat().st_size}
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            result["columns"] = next(reader, [])
            result["rows"] = sum(1 for _ in reader)
    elif path.suffix.lower() == ".xlsx":
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            result["sheets"] = [{"name": s.title, "rows": max(0, s.max_row - 1),
                                  "columns": [str(v) if v is not None else "" for v in next(s.values, [])]}
                                 for s in book]
        finally:
            book.close()
    else:
        result["warning"] = "Metadata only; inspect this attachment with its native reader"
    return result


def parse(path: Path, attachments: list[Path] | None = None) -> dict:
    pages, warnings = extract_pages(path)
    numbered = [(p, l, line) for p, text in enumerate(pages, 1)
                for l, line in enumerate(text.splitlines(), 1)]
    starts = [i for i, (_, _, line) in enumerate(numbered) if QUESTION.match(line)]
    questions = []
    for n, start in enumerate(starts):
        segment = numbered[start:starts[n+1] if n+1 < len(starts) else len(numbered)]
        raw = "\n".join(item[2] for item in segment)
        questions.append({"candidate_id": f"Q{n+1}", "heading": segment[0][2],
                          "source": {"page": segment[0][0], "line": segment[0][1]},
                          "raw_text": raw, "type_candidates": classify(raw),
                          "objective": None, "constraints": [], "deliverables": [],
                          "depends_on": [], "semantic_reviewed": False})
    info = []
    for item in attachments or []:
        try:
            info.append(attachment_info(item))
        except Exception as exc:
            info.append({"path": str(item), "error": str(exc)})
    if not starts:
        warnings.append("No reliable question headings found; agent must segment the extracted text")
    return {"schema_version": 1, "source_path": str(path.resolve()),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "pages": [{"page": n, "text": t} for n, t in enumerate(pages, 1)],
            "candidate_count": len(questions), "question_count": None,
            "questions": questions, "attachments": info, "warnings": warnings,
            "review_status": "needs_semantic_review"}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("problem", type=Path)
    ap.add_argument("--attachments", type=Path, nargs="*", default=[])
    ap.add_argument("--output", type=Path, required=True)
    args = ap.parse_args()
    result = parse(args.problem, args.attachments)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(args.output), "candidates": result["candidate_count"],
                      "warnings": result["warnings"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
