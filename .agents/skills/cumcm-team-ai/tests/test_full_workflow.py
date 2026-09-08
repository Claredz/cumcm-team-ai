"""Behavioral checks of the integrated pipeline, using explicitly synthetic inputs."""
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import corpus
import parse_problem
import pdf_audit
import prose_lint
import render_ai_usage
import render_paper
import workflow
import pymupdf


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        workflow.init(self.root, "cumcm", 2026)
        (self.root / "proof.txt").write_text("actual fixture evidence", encoding="utf-8")
        self.receipt = {"status": "passed", "checks": ["fixture checked"], "artifacts": ["proof.txt"], "issues": []}

    def test_resume_preserves_state_and_rejects_different_competition(self):
        workflow.complete(self.root, 0, self.receipt)
        self.assertEqual(workflow.init(self.root, "cumcm", 2026)["current_stage"], 1)
        with self.assertRaises(ValueError):
            workflow.init(self.root, "mcm", 2027)

    def test_out_of_order_completion_is_rejected(self):
        with self.assertRaises(ValueError):
            workflow.complete(self.root, 2, self.receipt)

    def test_missing_artifact_and_high_issue_block(self):
        bad = {**self.receipt, "artifacts": ["missing"]}
        with self.assertRaises(ValueError):
            workflow.complete(self.root, 0, bad)
        with self.assertRaises(ValueError):
            workflow.complete(self.root, 0, {**self.receipt, "issues": [{"severity": "high"}]})

    def test_changed_upstream_cannot_silently_advance(self):
        workflow.complete(self.root, 0, self.receipt)
        (self.root / "proof.txt").write_text("changed", encoding="utf-8")
        self.assertEqual(workflow.status(self.root)["changed_artifacts"], ["proof.txt"])
        with self.assertRaisesRegex(ValueError, "Upstream"):
            workflow.complete(self.root, 1, self.receipt)
        workflow.rollback(self.root, 0, "recheck source")
        self.assertEqual(workflow.status(self.root)["stale_stages"], [0])
        workflow.complete(self.root, 0, self.receipt)
        self.assertEqual(workflow.status(self.root)["changed_artifacts"], [])

    def test_formal_core_stage_requires_actual_human_review(self):
        state = workflow.load(self.root)
        state["workflow"]["formal_contest"] = True
        workflow.save(self.root / "state/decision_log.json", state)
        workflow.complete(self.root, 0, self.receipt)
        with self.assertRaisesRegex(ValueError, "human review"):
            workflow.complete(self.root, 1, self.receipt)

    def test_ten_stages_finish_only_with_compliance_evidence(self):
        for stage in range(9):
            workflow.complete(self.root, stage, self.receipt)
        with self.assertRaisesRegex(ValueError, "compliance"):
            workflow.complete(self.root, 9, self.receipt)
        state = workflow.load(self.root)
        state["stages"]["9"]["compliance_checks"] = {k: True for k in state["stages"]["9"]["compliance_checks"]}
        state["compliance"]["ai_usage"] = []  # synthetic fixture, not a real no-use declaration
        workflow.save(self.root / "state/decision_log.json", state)
        workflow.complete(self.root, 9, self.receipt)
        self.assertTrue(workflow.status(self.root)["finished"])


class ParsingTests(unittest.TestCase):
    def test_questions_source_positions_and_csv(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            p = root / "problem.md"
            p.write_text("背景资料\n问题一：预测需求。\n保留时间顺序。\n问题二：最小化成本。\n给出约束。", encoding="utf-8")
            a = root / "data.csv"
            a.write_text("time,value\n1,2\n2,3\n", encoding="utf-8-sig")
            result = parse_problem.parse(p, [a])
            self.assertEqual(result["candidate_count"], 2)
            self.assertIsNone(result["question_count"])
            self.assertEqual(result["questions"][0]["source"]["line"], 2)
            self.assertEqual(result["questions"][1]["type_candidates"][0]["type"], "optimization")
            self.assertEqual(result["attachments"][0]["rows"], 2)

    def test_docx_and_excel_are_read(self):
        from docx import Document
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            d = Document(); d.add_paragraph("问题一：评价排名。")
            d.save(root / "input.docx")
            b = openpyxl.Workbook(); b.active.append(["x", "y"]); b.active.append([1, 2]); b.save(root / "data.xlsx")
            report = parse_problem.parse(root / "input.docx", [root / "data.xlsx"])
            self.assertEqual(report["candidate_count"], 1)
            self.assertEqual(report["attachments"][0]["sheets"][0]["rows"], 1)

    def test_image_only_pdf_is_not_successfully_understood(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "blank.pdf"
            d = pymupdf.open(); d.new_page(); d.save(p); d.close()
            result = parse_problem.parse(p)
            self.assertTrue(result["warnings"])
            self.assertIsNone(result["question_count"])


class StyleTests(unittest.TestCase):
    def test_protected_values_cannot_be_changed_by_polishing(self):
        a = "误差从4.8降至4.1，结果不显著。$x=2$，见[1]。"
        b = "误差从4.8降至3.1，结果显著。$x=3$，见[2]。"
        self.assertEqual(set(prose_lint.compare(a, b)), {"numbers", "math", "citations", "negation"})

    def test_style_only_edits_preserve_protected_tokens(self):
        a = "值得注意的是，误差为 4.1，见[1]。"
        b = "误差为 4.1，见[1]。"
        self.assertTrue(prose_lint.lint(a))
        self.assertEqual(prose_lint.compare(a, b), {})

    def test_code_and_equations_are_not_style_linted(self):
        self.assertEqual(prose_lint.lint('```\n值得注意的是\n```\n$delve = 2$'), [])


class PdfCorpusTests(unittest.TestCase):
    @staticmethod
    def make_pdf(path, page_count, appendix=None, ai=None):
        doc = pymupdf.open()
        for i in range(1, page_count+1):
            p = doc.new_page()
            heading = "Appendix" if i == appendix else "Report on Use of AI" if i == ai else "Example page"
            p.insert_text((72, 72), heading)
            p.insert_textbox(pymupdf.Rect(72, 100, 500, 700), "Synthetic test text, not a real paper. " * 30)
        doc.save(path); doc.close()

    def test_cumcm_counts_body_not_appendix(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "paper.pdf"
            self.make_pdf(p, 34, appendix=32)
            result = pdf_audit.audit(p, "cumcm")
            self.assertEqual(result["counted_pages"], 30)
            self.assertFalse(any(i["code"] == "page-limit" for i in result["issues"]))

    def test_mcm_counts_appendix_but_excludes_ai_report(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "paper.pdf"
            self.make_pdf(p, 28, appendix=24, ai=27)
            result = pdf_audit.audit(p, "mcm")
            self.assertEqual(result["counted_pages"], 26)
            self.assertEqual(result["status"], "failed")

    def test_render_page_pngs(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "paper.pdf"
            self.make_pdf(p, 2)
            result = pdf_audit.audit(p, "mcm", render_dir=Path(tmp) / "pages")
            self.assertEqual(len(list((Path(tmp) / "pages").glob("*.png"))), 2)
            self.assertEqual(result["status"], "needs_visual_review")

    def test_shared_appendix_page_does_not_undercount_body(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "shared.pdf"
            d = pymupdf.open()
            for i in range(4):
                page = d.new_page()
                page.insert_text((72, 72), "Body text before any appendix heading.")
                if i == 3:
                    page.insert_text((72, 200), "Appendix")
            d.save(p); d.close()
            result = pdf_audit.audit(p, "diangong")
            self.assertEqual(result["counted_pages"], 2)
            self.assertTrue(any(x["code"] == "shared-boundary" for x in result["issues"]))

    def test_corpus_deduplicates_and_excludes_scans(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.make_pdf(root / "paper.pdf", 2)
            (root / "copy.pdf").write_bytes((root / "paper.pdf").read_bytes())
            d = pymupdf.open(); d.new_page(); d.save(root / "scan.pdf"); d.close()
            manifest = {"papers": [{"filename": n, "url": "https://example.org", "usage_basis": "self-authored fixture", "year": 2024, "topic": "A"} for n in ["paper.pdf", "copy.pdf", "scan.pdf"]]}
            index = corpus.ingest(root, manifest)
            self.assertEqual(sorted(r["status"] for r in index["records"]), ["accepted", "duplicate", "excluded"])
            self.assertEqual(corpus.stats(index)["accepted"], 1)
            self.assertEqual(corpus.stats({"records": []})["groups"]["all"]["dimensions"], {})


class DisclosureTests(unittest.TestCase):
    def test_used_ai_declaration_precedes_references(self):
        entry = {"tool": "Fixture AI", "model": "test", "version": "test", "use_stage": "8", "purpose": "语言润色", "paper_sections": ["摘要"], "disclosure": "synthetic test", "human_review": "fixture only"}
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); ws = root / "paper_workspace"; ws.mkdir()
            log = root / "decision_log.json"
            log.write_text(json.dumps({"competition": "cumcm", "compliance": {"ai_usage": [entry]}}), encoding="utf-8")
            render_ai_usage.render_reports(log, ws, root / "support", markdown_only=True)
            for section, name in render_paper.SECTION_TO_FILE.items():
                (ws / name).write_text("测试章节。", encoding="utf-8")
            tex, _ = render_paper.fill_template("cumcm", ws, root / "out", prefer_pandoc=False)
            body = tex.read_text(encoding="utf-8")
            self.assertLess(body.index(r"\input{sections/cumcm_no_ai_statement}"), body.index(r"\input{sections/8_references}"))
            self.assertIn("使用了AI工具", (ws / "AI工具使用声明.md").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
